"""问题2：数据审计、连续缺失增强、门控融合及蒸馏的可复现实验。

Only competition data train is used for fitting. Validation selects checkpoints.
The attachment-2 test split is evaluated once after selection; attachment 3 is
unlabelled inference only. Cache files are derived data and must not be committed.
"""
from __future__ import annotations
import argparse
import copy
import csv
import hashlib
import json
import os
import pickle
import random
import sys
import time
from pathlib import Path

import numpy as np
import torch
import torch.nn.functional as F
from torch import nn
from sklearn.metrics import accuracy_score, f1_score, mean_absolute_error, confusion_matrix

ROOT = Path(__file__).resolve().parents[1]
MODEL_ID = 'google-bert/bert-base-uncased'
MODEL_REVISION = '86b5e0934494bd15c9632b12f734a8a67f723594'
METHODS = {'baseline': '普通融合', 'augmentation': '连续缺失增强', 'gated': '动态门控', 'distilled': '门控蒸馏', 'hierarchical': '分层门控'}
MODS = 'TAV'


def dump(path, value):
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')


def write_csv(path, rows):
    path = Path(path); path.parent.mkdir(parents=True, exist_ok=True)
    if not rows: return
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)


def seed_all(seed):
    random.seed(seed); np.random.seed(seed); torch.manual_seed(seed)
    torch.cuda.manual_seed_all(seed)


def sha(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for b in iter(lambda: f.read(8 * 1024 * 1024), b''): h.update(b)
    return h.hexdigest()


def runtime(device):
    torch.set_num_threads(4)
    if device.startswith('cuda'):
        torch.cuda.set_device(device)
        torch.cuda.set_per_process_memory_fraction(.16, device)
    torch.backends.cudnn.benchmark = False
    torch.backends.cudnn.deterministic = True


def derive_masks(s):
    """Separate content timeline, text availability, A/V unobserved rows.

    Timeline is bounded by CLS/SEP, not attention.sum(): internal missing text
    must not shift subsequent positions. BERT PAD hidden states are not evidence.
    All-zero A/V within content is unavailable; cause cannot be inferred from zeros.
    """
    b = np.asarray(s['text_bert'])
    assert b.ndim == 3 and b.shape[1:] == (3, 50)
    assert np.isfinite(b).all() and np.equal(b, np.round(b)).all()
    b = b.astype(np.int64)
    assert np.isin(b[:, 1], [0, 1]).all()
    assert np.isin(b[:, 2], [0, 1]).all()
    assert b[:, 0].min() >= 0 and b[:, 0].max() < 30522
    n = len(b); valid = np.zeros((n, 50), bool)
    for i, row in enumerate(b):
        cls = np.flatnonzero(row[0] == 101); sep = np.flatnonzero(row[0] == 102)
        if not len(cls) or not len(sep):
            raise ValueError(f'样本{i}缺少CLS/SEP，不能可靠推断填充范围，需人工核对')
        valid[i, cls[0]+1:sep[-1]] = True
    # [UNK] 本身也是自然词元：完整版训练/验证中存在少量真正的 [UNK]。
    # 仅当它与音、视同位置全零同时出现时，才可判为附件3式共同缺失。
    audio_seen = np.any(np.asarray(s['audio']) != 0, axis=-1)
    vision_seen = np.any(np.asarray(s['vision']) != 0, axis=-1)
    joint_unknown = (b[:, 0] == 100) & ~audio_seen & ~vision_seen
    masks = np.stack([valid & (b[:, 1] > 0) & (b[:, 0] != 0) & ~joint_unknown,
                      valid & audio_seen, valid & vision_seen], -1)
    for m in ('audio', 'vision'):
        if not np.isfinite(s[m]).all():
            raise ValueError(f'{m}存在非有限值，先人工审计，不能静默填补')
    assert valid.any(1).all()
    return b, valid, masks


def audit(data, data_root, out):
    from openpyxl import load_workbook
    ws = load_workbook(data_root/'附件2-数据集特征文件/label.xlsx', read_only=True, data_only=True)['label']
    sheet = {f'{r[0]}$_${r[1]}': r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None}
    rows = []; summary = {}; errors = []
    for split, s in data.items():
        b, valid, masks = derive_masks(s)
        y = np.asarray(s['classification_labels']); r = np.asarray(s['regression_labels'])
        ids = list(map(str, s['id']))
        if len(set(ids)) != len(ids): errors.append(f'{split}:重复ID')
        if not np.isfinite(r).all() or np.any(np.abs(r)>3) or not np.array_equal(y, np.sign(r)+1):
            errors.append(f'{split}:标签越界或三分类与强度符号冲突')
        text_mismatch = 0
        for i, sid in enumerate(ids):
            original = sheet.get(sid)
            if original is None or float(original[3]) != float(r[i]) or str(original[5]) != split:
                errors.append(f'{sid}:标签表或划分不一致')
            if original and str(original[2]).strip() != str(s['raw_text'][i]).strip(): text_mismatch += 1
            L = int(valid[i].sum())
            rows.append({'划分': split, '样本编号': sid, '有效内容位置数': L,
                         '文本不可用位置': int(L-masks[i,:,0].sum()),
                         '语音不可用位置': int(L-masks[i,:,1].sum()),
                         '视觉不可用位置': int(L-masks[i,:,2].sum()),
                         '达到50位置截断上限': bool(b[i,1].sum()==50),
                         '处理': '保留；用模态掩码表示不可用位置',
                         '文本预览': str(s['raw_text'][i])[:160]})
        summary[split] = {'samples':len(ids), 'class_counts':np.bincount(y.astype(int), minlength=3).tolist(),
            'native_all_audio_unavailable':int((masks[:,:,1].sum(1)==0).sum()),
            'native_all_vision_unavailable':int((masks[:,:,2].sum(1)==0).sum()),
            'truncated_at_50':int((b[:,1].sum(1)==50).sum()), 'label_table_text_mismatches':text_mismatch,
            'content_positions':int(valid.sum()), 'observed_positions':masks.sum((0,1)).tolist()}
    for a, b in [('train','valid'), ('train','test'), ('valid','test')]:
        ia=set(map(str,data[a]['id'])); ib=set(map(str,data[b]['id']))
        va={x.split('$_$')[0] for x in ia}; vb={x.split('$_$')[0] for x in ib}
        summary[f'{a}_{b}_overlap']={'sample_ids':len(ia&ib),'video_ids':len(va&vb)}
        if ia&ib: errors.append(f'{a}/{b}:重复样本')
    summary['errors']=errors
    summary['rule']='不删除困难样本；不按模型预测修改标签；不将CLS/SEP/PAD计入缺失率；全零仅表示不可观测，不能推断人为缺失原因。'
    write_csv(out/'问题2_样本审计清单.csv', rows)
    write_csv(out/'问题2_需复核样本.csv', [r for r in rows if r['视觉不可用位置']==r['有效内容位置数'] or r['语音不可用位置']==r['有效内容位置数'] or r['达到50位置截断上限']])
    dump(out/'问题2_数据审计.json', summary)
    if errors: raise ValueError(errors[:10])
    return summary


class TextEncoder:
    def __init__(self, device, revision=None):
        from transformers import AutoModel
        self.device=device
        self.model=AutoModel.from_pretrained(MODEL_ID, revision=revision or MODEL_REVISION, local_files_only=True).to(device).eval()
        for p in self.model.parameters(): p.requires_grad_(False)
        self.revision=getattr(self.model.config, '_commit_hash', None)

    @torch.inference_mode()
    def encode(self, b, observed, batch=64):
        result=[]
        for start in range(0,len(b),batch):
            inp=torch.as_tensor(b[start:start+batch].copy(),device=self.device).long()
            obs=torch.as_tensor(observed[start:start+batch],device=self.device)
            special=(inp[:,0]==101)|(inp[:,0]==102)
            # Redact BEFORE contextual BERT encoding, preventing hidden tokens
            # leaking through surrounding token representations.
            redact=(~obs)&(~special)
            inp[:,0][redact]=103; inp[:,1][redact]=0
            z=self.model(input_ids=inp[:,0],attention_mask=inp[:,1],token_type_ids=inp[:,2]).last_hidden_state
            z=z*obs.unsqueeze(-1)
            result.append(z.cpu().numpy().astype(np.float16))
        return np.concatenate(result)


def block_mask(valid, rate, seed, position='random'):
    rng=np.random.default_rng(seed); hide=np.zeros_like(valid)
    for i in range(len(valid)):
        idx=np.flatnonzero(valid[i]); L=len(idx)
        k=max(1,min(L,int(round(rate*L)))) if rate else 0
        if not k: continue
        if position=='start': st=0
        elif position=='middle': st=(L-k)//2
        elif position=='end': st=L-k
        else: st=int(rng.integers(0,L-k+1))
        hide[i,idx[st:st+k]]=True
    return hide


def text_features(enc, b, masks, cache, name, source_hash):
    cache.mkdir(parents=True,exist_ok=True)
    fingerprint=hashlib.sha256((source_hash+str(enc.revision)+'q2_text_v1').encode()+b.tobytes()+masks.tobytes()).hexdigest()[:20]
    path=cache/f'{name}_{fingerprint}.npy'
    if path.exists(): return np.load(path)
    z=enc.encode(b,masks); np.save(path,z); return z


def prepare_split(s, enc, cache, name, source_hash):
    b,v,m=derive_masks(s)
    return {'b':b,'valid':v,'mask':m,'t':text_features(enc,b,m[:,:,0],cache,name,source_hash),
        'a':np.asarray(s['audio'],np.float32),'v':np.asarray(s['vision'],np.float32),
        'y':np.asarray(s.get('classification_labels',[]),np.int64),
        'r':np.asarray(s.get('regression_labels',[]),np.float32),
        'ids':list(map(str,s.get('id',[f'{name}_{i}' for i in range(len(b))]))),
        'raw':list(map(str,s.get('raw_text',['']*len(b))))}


def fit_normalization(tr):
    stats={}
    for j,k in enumerate(('t','a','v')):
        x=tr[k][tr['mask'][:,:,j]].astype(np.float32)
        stats[k]={'mean':x.mean(0).tolist(),'std':np.maximum(x.std(0),1e-4).tolist()}
    return stats


def normalize(s, stats):
    for j,k in enumerate(('t','a','v')):
        mean=np.asarray(stats[k]['mean'],np.float32); std=np.asarray(stats[k]['std'],np.float32)
        s[k]=np.clip((s[k].astype(np.float32)-mean)/std,-8,8)*s['mask'][:,:,j,None]
    return s


def norm_text(x, observed, stats):
    return np.clip((x.astype(np.float32)-np.asarray(stats['t']['mean'],np.float32))/np.asarray(stats['t']['std'],np.float32),-8,8)*observed[:,:,None]


class Fusion(nn.Module):
    def __init__(self, gated=False, hidden=96, dropout=.2):
        super().__init__(); self.gated=gated; self.hidden=hidden; self.dropout=dropout
        self.projections=nn.ModuleList([nn.Sequential(nn.Linear(d,hidden),nn.LayerNorm(hidden),nn.GELU()) for d in (768,74,35)])
        self.gate=nn.Sequential(nn.Linear(3*hidden+6,64),nn.GELU(),nn.Linear(64,3))
        self.fuse=nn.Sequential(nn.Linear(3*hidden+6,hidden),nn.GELU(),nn.Dropout(dropout))
        self.position=nn.Parameter(torch.randn(1,50,hidden)*.01)
        layer=nn.TransformerEncoderLayer(hidden,4,hidden*2,dropout,batch_first=True,norm_first=True)
        self.temporal=nn.TransformerEncoder(layer,1,enable_nested_tensor=False)
        self.attention=nn.Linear(hidden,1)
        self.head=nn.Sequential(nn.LayerNorm(hidden),nn.Linear(hidden,64),nn.GELU(),nn.Dropout(dropout))
        self.cls=nn.Linear(64,3); self.reg=nn.Linear(64,1)

    def forward(self,t,a,v,mask,valid):
        h=torch.stack([proj(x)*mask[:,:,j,None] for j,(proj,x) in enumerate(zip(self.projections,(t,a,v)))],dim=2)
        frac=mask.float().sum(1)/valid.sum(1,keepdim=True).clamp_min(1)
        quality=torch.cat([mask.float(),frac[:,None].expand(-1,50,-1)],-1)
        features=torch.cat([h.flatten(2),quality],-1)
        scores=self.gate(features) if self.gated else torch.zeros_like(mask,dtype=t.dtype)
        weights=torch.softmax(scores.masked_fill(~mask,-1e4),-1)*mask
        weights=weights/weights.sum(-1,keepdim=True).clamp_min(1e-8)
        z=self.fuse(torch.cat([(h*weights.unsqueeze(-1)).flatten(2),quality],-1))+self.position
        z=self.temporal(z,src_key_padding_mask=~valid)
        observed=valid & mask.any(-1)
        w=torch.softmax(self.attention(z).squeeze(-1).masked_fill(~observed,-1e4),1)*observed
        w=w/w.sum(1,keepdim=True).clamp_min(1e-8)
        pooled=(z*w.unsqueeze(-1)).sum(1)
        pooled=self.head(pooled)
        return self.cls(pooled),3*torch.tanh(self.reg(pooled).squeeze(-1)/3)


class HierarchicalFusion(Fusion):
    """Two decisions: neutral vs polar, then polarity among polar examples."""
    def __init__(self, hidden=96, dropout=.2):
        super().__init__(gated=True, hidden=hidden, dropout=dropout)
        self.cls=nn.Linear(64,2)

    def forward(self,t,a,v,mask,valid):
        logits,intensity=super().forward(t,a,v,mask,valid)
        neutral=torch.sigmoid(logits[:,0]); positive=torch.sigmoid(logits[:,1])
        probability=torch.stack([(1-neutral)*(1-positive),neutral,(1-neutral)*positive],-1)
        return torch.log(probability.clamp_min(1e-8)),intensity


def metrics(y,r,p,pr):
    pred=p.argmax(1)
    corr=float(np.corrcoef(r,pr)[0,1]) if np.std(pr)>1e-8 and np.std(r)>1e-8 else 0.
    return {'Accuracy':float(accuracy_score(y,pred)), 'Macro_F1':float(f1_score(y,pred,average='macro',zero_division=0)),
        'Weighted_F1':float(f1_score(y,pred,average='weighted',zero_division=0)),
        'MAE':float(mean_absolute_error(r,pr)), 'Pearson':corr}


def tensor_batch(s,idx,device,override_t=None,override_mask=None):
    m=s['mask'][idx] if override_mask is None else override_mask
    xs=[s['t'][idx] if override_t is None else override_t,s['a'][idx],s['v'][idx]]
    return [torch.as_tensor(x*m[:,:,j,None],dtype=torch.float32,device=device) for j,x in enumerate(xs)]+[
        torch.as_tensor(m,device=device),torch.as_tensor(s['valid'][idx],device=device)]


@torch.inference_mode()
def predict(model,s,device,scenario=None):
    model.eval(); probs=[]; regs=[]
    for start in range(0,len(s['ids']),128):
        idx=np.arange(start,min(start+128,len(s['ids'])))
        t=scenario['t'][idx] if scenario and scenario.get('t') is not None else None
        m=scenario['mask'][idx] if scenario else None
        lg,r=model(*tensor_batch(s,idx,device,t,m)); probs.append(lg.softmax(-1).cpu().numpy()); regs.append(r.cpu().numpy())
    return np.concatenate(probs),np.concatenate(regs)


def scenarios(s,enc,cache,stats,source_hash,full=False,repeats=2):
    specs=[('none',0.,'random',0)]
    if full:
        for sub in ('T','A','V','TA','TV','AV','TAV'):
            for rate in (.1,.3,.5,.7):
                for rep in range(repeats): specs.append((sub,rate,'random',rep))
        for sub in ('T','A','V','AV'):
            for pos in ('start','middle','end'): specs.append((sub,.3,pos,0))
    else:
        for sub in ('T','A','V','AV','TAV'): specs.append((sub,.3,'random',0))
    banks={}; output=[]
    for sub,rate,pos,rep in specs:
        mask=s['mask'].copy(); text=None
        if rate:
            seed=82000+rep*1009+int(rate*100)
            # Matched gap locations across methods; same temporal gap across modalities.
            hide=block_mask(s['valid'],rate,seed,pos)
            for ch in sub: mask[:,:,MODS.index(ch)] &= ~hide
            if 'T' in sub:
                key=(rate,pos,rep)
                if key not in banks:
                    z=text_features(enc,s['b'],mask[:,:,0],cache,'evaluation_text',source_hash)
                    banks[key]=norm_text(z,mask[:,:,0],stats)
                text=banks[key]
        output.append({'subset':sub,'rate':rate,'position':pos,'replicate':rep,'mask':mask,'t':text})
    return output


def evaluate(model,s,suite,device):
    rows=[]; allpred=[]
    for sc in suite:
        p,pr=predict(model,s,device,sc)
        row={k:sc[k] for k in ('subset','rate','position','replicate')}
        row.update(metrics(s['y'],s['r'],p,pr))
        for j,ch in enumerate(MODS):
            row[f'{ch}_实际不可用比例']=float((s['valid'].sum()-sc['mask'][:,:,j].sum())/s['valid'].sum())
        rows.append(row); allpred.append((p,pr))
    return rows,allpred


def train_one(method,seed,tr,va,quick,banks,device,out,epochs,teacher,stats,provenance):
    seed_all(seed); model=(HierarchicalFusion() if method=='hierarchical' else Fusion(gated=method in ('gated','distilled'))).to(device)
    opt=torch.optim.AdamW(model.parameters(),lr=3e-4,weight_decay=1e-3)
    weights=torch.as_tensor(np.sqrt(len(tr['y'])/(3*np.bincount(tr['y'],minlength=3))),dtype=torch.float32,device=device)
    teacher_pred=predict(teacher,tr,device) if teacher is not None and method=='distilled' else None
    best=-float('inf'); bad=0; logs=[]; checkpoint=out/f'问题2_{METHODS[method]}_种子{seed}.pt'
    for ep in range(1,epochs+1):
        model.train(); permutation=np.random.permutation(len(tr['ids'])); losses=[]
        for st in range(0,len(permutation),64):
            idx=permutation[st:st+64]; mask=tr['mask'][idx].copy(); text=tr['t'][idx].copy()
            changed=np.zeros(len(idx),bool)
            if method!='baseline':
                for k,i in enumerate(idx):
                    if np.random.rand()>.65: continue
                    sub=random.choice(('T','A','V','TA','TV','AV','TAV'))
                    bank=random.randrange(len(banks)); hide=banks[bank]['hide'][i]
                    for ch in sub: mask[k,:,MODS.index(ch)] &= ~hide
                    if 'T' in sub: text[k]=banks[bank]['t'][i]
                    changed[k]=True
            lg,pr=model(*tensor_batch(tr,idx,device,text,mask))
            y=torch.as_tensor(tr['y'][idx],device=device); r=torch.as_tensor(tr['r'][idx],device=device)
            loss=F.cross_entropy(lg,y,weight=weights)+.6*F.smooth_l1_loss(pr,r)
            if teacher_pred is not None:
                tp,trr=teacher_pred
                eligible=changed & (tp[idx].argmax(1)==tr['y'][idx]) & (np.abs(trr[idx]-tr['r'][idx])<1.)
                if eligible.any():
                    tm=torch.as_tensor(eligible,device=device)
                    probs=torch.as_tensor(tp[idx][eligible],device=device)
                    soft=F.softmax(torch.log(probs.clamp_min(1e-8))/2,-1)
                    loss=loss+.2*4*F.kl_div(F.log_softmax(lg[tm]/2,-1),soft,reduction='batchmean')
                    loss=loss+.1*F.smooth_l1_loss(pr[tm],torch.as_tensor(trr[idx][eligible],device=device))
            opt.zero_grad(set_to_none=True); loss.backward(); nn.utils.clip_grad_norm_(model.parameters(),1.); opt.step(); losses.append(float(loss.detach()))
        rows,_=evaluate(model,va,quick,device)
        clean=rows[0]; robust={k:float(np.mean([x[k] for x in rows[1:]])) for k in ('Macro_F1','MAE')}
        score=.5*clean['Macro_F1']+.5*robust['Macro_F1']-.1*(clean['MAE']+robust['MAE'])/2
        logs.append({'模型':METHODS[method],'种子':seed,'轮次':ep,'训练损失':float(np.mean(losses)),
            '验证完整F1':clean['Macro_F1'],'验证完整MAE':clean['MAE'],'验证缺失F1':robust['Macro_F1'],'验证缺失MAE':robust['MAE'],'选模分数':score})
        print(json.dumps(logs[-1],ensure_ascii=False),flush=True)
        if score>best+1e-4:
            best=score; bad=0
            torch.save({'model':model.state_dict(),'method':method,'seed':seed,'epoch':ep,'selection_score':score,
                'stats':stats,'provenance':provenance,'architecture':{'hidden':96,'dropout':.2},'class_names':['Negative','Neutral','Positive']},checkpoint)
        else: bad+=1
        write_csv(out/f'问题2_{METHODS[method]}_种子{seed}_训练记录.csv',logs)
        if bad>=7: break
    model.load_state_dict(torch.load(checkpoint,map_location=device,weights_only=False)['model']); model.eval()
    return model,checkpoint,logs


def consistency_checks(model,s,device):
    """Check actual failure risks: ignored padding, unavailable signals, empty inputs."""
    idx=np.arange(min(8,len(s['ids']))); model.eval()
    args=tensor_batch(s,idx,device)
    with torch.inference_mode():
        a=model(*args)
        perturbed=[x.clone() for x in args]
        for j in range(3): perturbed[j][~args[3][:,:,j]]=12345.
        b=model(*perturbed)
        assert all(torch.allclose(x,y,atol=1e-5) for x,y in zip(a,b)), '缺失/填充输入影响了预测'
        empty=[x.clone() for x in args]; empty[3].zero_()
        z=model(*empty)
        assert all(torch.isfinite(x).all() for x in z), '全缺失产生非有限值'
        assert (z[1].abs()<=3).all()
    return {'missing_value_invariance':True,'all_missing_finite':True,'intensity_in_range':True}


def save_predictions(out,stem,s,p,pr):
    names=['Negative','Neutral','Positive']; rows=[]
    for i,sid in enumerate(s['ids']):
        row={'样本编号':sid,'预测极性':names[int(p[i].argmax())],'预测强度':float(pr[i]),'预测最大概率':float(p[i].max()),
            '负向概率':float(p[i,0]),'中性概率':float(p[i,1]),'正向概率':float(p[i,2])}
        if len(s['y']): row.update({'真实极性':names[int(s['y'][i])],'真实强度':float(s['r'][i]),'分类是否错误':bool(p[i].argmax()!=s['y'][i]),'绝对误差':float(abs(pr[i]-s['r'][i]))})
        L=s['valid'][i].sum()
        for j,ch in enumerate(('文本','语音','视觉')): row[ch+'不可用比例']=float(1-s['mask'][i,:,j].sum()/L)
        row['文本']=s['raw'][i]; rows.append(row)
    write_csv(out/stem,rows)
    return rows


def run(args):
    runtime(args.device); seed_all(args.seed)
    root=Path(args.data_root); out=Path(args.out); out.mkdir(parents=True,exist_ok=True)
    cache=ROOT/'work/问题2_特征缓存'; source=root/'附件2-数据集特征文件/aligned_50.pkl'
    started=time.time(); source_hash=sha(source)
    with source.open('rb') as f: data=pickle.load(f)
    audit(data,root,out)
    encoder=TextEncoder(args.device)
    provenance={'data_sha256':source_hash,'code_sha256':sha(__file__),'bert_model':MODEL_ID,'bert_revision':encoder.revision,
        'python':sys.version.split()[0],'torch':torch.__version__,'numpy':np.__version__, 'device':torch.cuda.get_device_name(args.device) if args.device.startswith('cuda') else 'cpu',
        'split_sizes':{k:len(v['id']) for k,v in data.items()},'max_epochs':args.epochs,'batch_size':64,
        'learning_rate':3e-4,'weight_decay':1e-3,'hidden':96,'dropout':.2,'patience':7,'gradient_clip':1.,
        'augmentation_probability':.65,'augmentation_rates':[.1,.3,.5],'missing_before_BERT':True,
        'classification_loss':'sqrt inverse frequency weighted CE','regression_loss':'0.6 SmoothL1',
        'KD':'0.2 KL temperature=2 + 0.1 SmoothL1; only correct teacher train predictions with abs error <1',
        'selection':'0.5 clean MacroF1 + 0.5 mean missing MacroF1 - 0.1 mean(clean MAE, missing MAE)',
        'validation_selection_scenarios':'clean and T/A/V/AV/TAV 30% contiguous random gaps (fixed seed)',
        'test_policy':'test performance not used for fitting, early stopping or candidate selection'}
    dump(out/'问题2_实验配置.json',provenance)
    tr=prepare_split(data['train'],encoder,cache,'train',source_hash)
    va=prepare_split(data['valid'],encoder,cache,'valid',source_hash)
    stats=fit_normalization(tr); normalize(tr,stats); normalize(va,stats)
    banks=[]
    for rate in (.1,.3,.5):
        hide=block_mask(tr['valid'],rate,51000+int(rate*100))
        observed=tr['mask'][:,:,0]&~hide
        text=text_features(encoder,tr['b'],observed,cache,'train_aug',source_hash)
        banks.append({'hide':hide,'t':norm_text(text,observed,stats)})
        print(f'prepared training text gaps {rate}',flush=True)
    quick=scenarios(va,encoder,cache,stats,source_hash)
    models={}; paths={}; allrows=[]
    for method in ('baseline','augmentation','gated','distilled'):
        model,path,_=train_one(method,args.seed,tr,va,quick,banks,args.device,out,args.epochs,models.get('baseline'),stats,provenance)
        models[method]=model; paths[method]=path
    suite=scenarios(va,encoder,cache,stats,source_hash,full=True,repeats=2)
    val_results={}; val_preds={}
    for method,model in models.items():
        checks=consistency_checks(model,va,args.device)
        rows,preds=evaluate(model,va,suite,args.device); val_results[method]=rows; val_preds[method]=preds
        allrows.extend([{'模型':METHODS[method],'种子':args.seed,**r} for r in rows])
        save_predictions(out,f'问题2_验证集_{METHODS[method]}_逐条预测.csv',va,*preds[0])
    write_csv(out/'问题2_验证集缺失实验.csv',allrows)
    # Candidate selection uses validation only and the same predefined objective.
    def score(rows):
        c=rows[0]; r=[x for x in rows[1:] if x['position']=='random']
        return .5*c['Macro_F1']+.5*np.mean([x['Macro_F1'] for x in r])-.1*(c['MAE']+np.mean([x['MAE'] for x in r]))/2
    selected=max(('augmentation','gated','distilled'),key=lambda m:score(val_results[m]))
    selected_info={'selected_method':selected,'selected_name':METHODS[selected],
        'validation_scores':{m:float(score(rows)) for m,rows in val_results.items()},'model_file':paths[selected].name,
        'meaningful_checks':checks,'elapsed_sec_before_test':time.time()-started}
    dump(out/'问题2_验证集选模结论.json',selected_info)
    # Held-out evaluation is deliberately after selection and must not feed back.
    te=normalize(prepare_split(data['test'],encoder,cache,'test',source_hash),stats)
    tsuite=scenarios(te,encoder,cache,stats,source_hash,full=False)
    testrows=[]
    for method in ('baseline',selected):
        rows,preds=evaluate(models[method],te,tsuite,args.device)
        testrows.extend([{'模型':METHODS[method],'种子':args.seed,**r} for r in rows])
        save_predictions(out,f'问题2_留出测试集_{METHODS[method]}_逐条预测.csv',te,*preds[0])
    write_csv(out/'问题2_留出测试集一次性评价.csv',testrows)
    # Final unlabelled predictions. No label lookup or matching to labelled samples.
    all_test=[]
    for file in sorted((root/'附件3-模态缺失特征样本/对齐版本').glob('*.pkl')):
        with file.open('rb') as f: s=pickle.load(f)['test']
        s['id']=[file.stem] if len(s['audio'])==1 else [f'{file.stem}_{i}' for i in range(len(s['audio']))]
        item=normalize(prepare_split(s,encoder,cache,'annex3',sha(file)),stats)
        p,r=predict(models[selected],item,args.device)
        rows=save_predictions(out/'逐样本预测',f'{file.stem}.csv',item,p,r)
        for row in rows: row['所用模型']=METHODS[selected]
        all_test.extend(rows)
    assert len(all_test)==30 and len({r['样本编号'] for r in all_test})==30
    write_csv(out/'问题2_附件3全量预测.csv',all_test)
    dump(out/'问题2_完成状态.json',{'complete':True,'annex3_predictions':len(all_test),'selected':selected,
         'total_elapsed_sec':time.time()-started,'tests':checks,'git_push_performed':False})
    print('FINISHED',json.dumps(selected_info,ensure_ascii=False),flush=True)


if __name__=='__main__':
    p=argparse.ArgumentParser(); p.add_argument('--data-root',default=str(ROOT/'E题数据/E题数据'))
    p.add_argument('--out',default=str(ROOT/'outputs/问题2_首轮实验结果'))
    p.add_argument('--device',default='cuda:1'); p.add_argument('--epochs',type=int,default=30); p.add_argument('--seed',type=int,default=2026)
    run(p.parse_args())
