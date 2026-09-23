"""Independent full-MOSEI train/validation experiment for Problem 2.

Fit only the 16,326 official full-data train segments. Select on its official
1,871-segment valid split. Original attachment-2 test and attachment 3 are
never used for fitting or selection; their IDs are consulted only for leakage
checks. No existing Q2 checkpoint or result is overwritten.
"""
from __future__ import annotations

import argparse
import csv
import gc
import json
import pickle
import random
import sys
import time
from pathlib import Path

import numpy as np
import torch
import torch.nn.functional as F
from openpyxl import load_workbook
from torch import nn

from q2_experiment import (
    ROOT, Fusion, TextEncoder, derive_masks, dump, fit_normalization, metrics,
    normalize, predict, prepare_split, runtime, scenarios, seed_all, sha,
    tensor_batch, write_csv,
)
from q2_scattered_gap_experiment import (
    augment_batch, get_banks, make_scattered_suite, profile,
)


def video_id(sample_id):
    stem, sep, clip = str(sample_id).partition('$_$')
    if not sep or not stem or not clip:
        raise ValueError(f'Invalid MOSEI sample ID: {sample_id!r}')
    return stem


def read_target_ids(data_root):
    sheet = load_workbook(data_root / '附件2-数据集特征文件/label.xlsx',
                          read_only=True, data_only=True)['label']
    splits = {'train': set(), 'valid': set(), 'test': set()}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        split = str(row[5])
        if split not in splits:
            raise ValueError(f'Unexpected target split: {split}')
        splits[split].add(f'{row[0]}$_${row[1]}')
    return splits


def label_csv_audit(csv_path, splits):
    if not csv_path.exists():
        raise FileNotFoundError(f'Full-data label.csv missing: {csv_path}')
    listed = {'train': {}, 'valid': {}}
    with csv_path.open('r', encoding='utf-8-sig', newline='') as stream:
        for row in csv.DictReader(stream):
            mode = row['mode']
            if mode == 'test':
                # Do not inspect the full-test identifiers or labels: Annex 3
                # contains anonymous samples sourced from that split.
                continue
            if mode not in listed:
                raise ValueError(f'Unexpected full-data mode: {mode}')
            sample_id = f"{row['video_id']}$_${row['clip_id']}"
            if sample_id in listed[mode]:
                raise ValueError(f'Duplicate label.csv ID: {sample_id}')
            listed[mode][sample_id] = row
    for name, split in splits.items():
        ids = list(map(str, split['id']))
        if set(ids) != set(listed[name]):
            raise ValueError(f'{name}: aligned pickle IDs differ from label.csv')
        intensity = np.asarray(split['regression_labels'], dtype=np.float64)
        classification = np.asarray(split['classification_labels'], dtype=np.int64)
        if not np.isfinite(intensity).all() or not np.array_equal(
                classification, np.sign(intensity).astype(np.int64) + 1):
            raise ValueError(f'{name}: classification/regression label mismatch')
        if np.abs(intensity).max() > 3:
            raise ValueError(f'{name}: sentiment intensity outside [-3,3]')
        for sample_id, reg in zip(ids, intensity):
            if not np.isclose(reg, float(listed[name][sample_id]['label']), atol=1e-6):
                raise ValueError(f'{sample_id}: pickle/CSV regression mismatch')
    return {name: {'samples': len(rows),
                   'videos': len({video_id(sample_id) for sample_id in rows})}
            for name, rows in listed.items()}


def audit_splits(data, data_root, csv_path, safety_report_path):
    required = {'raw_text', 'audio', 'vision', 'id', 'text_bert',
                'classification_labels', 'regression_labels'}
    for name in ('train', 'valid'):
        split = data[name]
        if not required.issubset(split):
            raise ValueError(f'{name} missing fields {required-set(split)}')
        count = len(split['id'])
        if np.asarray(split['text_bert']).shape != (count, 3, 50):
            raise ValueError(f'{name}: BERT token shape mismatch')
        if np.asarray(split['audio']).shape != (count, 50, 74):
            raise ValueError(f'{name}: audio shape mismatch')
        if np.asarray(split['vision']).shape != (count, 50, 35):
            raise ValueError(f'{name}: vision shape mismatch')
        if len(set(map(str, split['id']))) != count:
            raise ValueError(f'{name}: duplicated sample IDs')
        derive_masks(split)
    csv_counts = label_csv_audit(csv_path,
                                 {name:data[name] for name in ('train','valid')})
    ids = {name: set(map(str, data[name]['id'])) for name in ('train','valid')}
    videos = {name: {video_id(x) for x in value} for name, value in ids.items()}
    if ids['train'] & ids['valid'] or videos['train'] & videos['valid']:
        raise ValueError('Full-data train/valid sample or video leakage')
    target = read_target_ids(data_root)
    for name in ('train','valid'):
        if not target[name].issubset(ids[name]):
            raise ValueError(f'Attachment-2 {name} IDs not contained in matching full split')
    protected = {video_id(x) for x in target['valid'] | target['test']}
    if videos['train'] & protected:
        raise ValueError('Full-data train overlaps protected attachment-2 videos')
    if len(ids['train']) != 16326 or len(ids['valid']) != 1871:
        raise ValueError('Unexpected full-data train/valid sizes; review source version')
    with safety_report_path.open('r', encoding='utf-8') as stream:
        safety = json.load(stream)
    if safety.get('annex3_unique_matches') != 30 or safety.get('annex3_match_split_counts') != {'test':30}:
        raise ValueError('Independent Annex-3 leakage audit does not clear full train/valid')
    if safety.get('train_videos_excluded_due_to_annex3_overlap') or safety.get('safe_train_sample_count_after_annex3_exclusion') != len(ids['train']):
        raise ValueError('Full train includes Annex-3 overlap; stop and exclude groups')
    summary = {'full_data': csv_counts,
               'attachment2': {name: len(value) for name, value in target.items()},
               'full_train_additional_to_attachment2_train': len(ids['train']-target['train']),
               'full_train_vs_attachment2_valid_test_video_overlap': 0,
               'full_train_valid_video_overlap': 0,
               'annex3_train_valid_feature_matches': 0,
               'independent_annex3_safety_audit_sha256': sha(safety_report_path),
               'label_mapping': 'Negative=0 for r<0; Neutral=1 for r=0; Positive=2 for r>0',
               'policy': 'Only full-data train used to fit; full-data valid selects. Full/attachment2 test excluded.'}
    return summary, target


def hide_joint_unknown(split):
    """An [UNK] at a simultaneous A/V zero row is missing, not speech content."""
    bert = np.asarray(split['text_bert'])
    in_content = np.asarray(split['audio']).shape[1] == 50
    if not in_content:
        raise ValueError('Unexpected alignment length')
    joint_zero = (np.asarray(split['audio']) == 0).all(-1) & (
        np.asarray(split['vision']) == 0).all(-1)
    affected = (bert[:, 0] == 100) & (bert[:, 1] > 0) & joint_zero
    count = int(affected.sum())
    if count:
        bert = bert.copy()
        bert[:, 1][affected] = 0
        split['text_bert'] = bert
    return count


def class_weights(train, device):
    counts = np.bincount(train['y'], minlength=3)
    if (counts == 0).any():
        raise ValueError('Class absent from full-data train')
    return torch.as_tensor(np.sqrt(len(train['y']) / (3 * counts)),
                           dtype=torch.float32, device=device)


def save_checkpoint(path, model, method, seed, epoch, selection_score, stats, config):
    torch.save({'model': model.state_dict(), 'method': method, 'seed': seed,
                'epoch': epoch, 'selection_score': float(selection_score),
                'stats': stats, 'provenance': config,
                'architecture': {'hidden': 96, 'dropout': .2}}, path)


def train_teacher(train, valid, quick, device, out, config, stats, seed, epochs, patience):
    model = Fusion(gated=False).to(device)
    optimizer = torch.optim.AdamW(model.parameters(), lr=3e-4, weight_decay=1e-3)
    weights = class_weights(train, device)
    checkpoint = out / f'完整数据_干净教师_种子{seed}.pt'
    best, stale, history = -float('inf'), 0, []
    for epoch in range(1, epochs+1):
        model.train()
        order = np.random.permutation(len(train['ids']))
        losses = []
        for start in range(0, len(order), 64):
            idx = order[start:start+64]
            logits, reg = model(*tensor_batch(train, idx, device))
            y = torch.as_tensor(train['y'][idx], device=device)
            r = torch.as_tensor(train['r'][idx], device=device)
            loss = F.cross_entropy(logits, y, weight=weights) + .6 * F.smooth_l1_loss(reg, r)
            optimizer.zero_grad(set_to_none=True)
            loss.backward()
            nn.utils.clip_grad_norm_(model.parameters(), 1.)
            optimizer.step()
            losses.append(float(loss.detach()))
        probability, intensity = predict(model, valid, device, quick[0])
        clean = metrics(valid['y'], valid['r'], probability, intensity)
        selection = clean['Macro_F1'] - .1 * clean['MAE']
        row = {'stage':'teacher', 'epoch':epoch, 'train_loss':float(np.mean(losses)),
               'clean_F1':clean['Macro_F1'], 'clean_MAE':clean['MAE'],
               'selection_score':selection}
        history.append(row)
        print(json.dumps(row, ensure_ascii=False), flush=True)
        write_csv(out / '完整数据_干净教师_训练记录.csv', history)
        if selection > best + 1e-4:
            best, stale = selection, 0
            save_checkpoint(checkpoint, model, 'baseline', seed, epoch, selection, stats, config)
        else:
            stale += 1
        if stale >= patience:
            break
    model.load_state_dict(torch.load(checkpoint, map_location=device, weights_only=False)['model'])
    model.eval()
    return model, checkpoint


def train_student(train, valid, quick, sparse, teacher, banks, device, out,
                  config, stats, seed, epochs, patience):
    blocks, scattered = banks
    teacher_prob, teacher_reg = predict(teacher, train, device)
    teacher_eligible = (teacher_prob.argmax(1) == train['y']) & (
        np.abs(teacher_reg-train['r']) < 1.)
    model = Fusion(gated=True).to(device)
    optimizer = torch.optim.AdamW(model.parameters(), lr=3e-4, weight_decay=1e-3)
    weights = class_weights(train, device)
    checkpoint = out / f'完整数据_缺失蒸馏学生_种子{seed}.pt'
    best, stale, history = -float('inf'), 0, []
    for epoch in range(1, epochs+1):
        model.train()
        order = np.random.permutation(len(train['ids']))
        losses = []
        for start in range(0, len(order), 64):
            idx = order[start:start+64]
            mask, text, changed = augment_batch(train, idx, blocks, scattered)
            logits, reg = model(*tensor_batch(train, idx, device, text, mask))
            y = torch.as_tensor(train['y'][idx], device=device)
            r = torch.as_tensor(train['r'][idx], device=device)
            loss = F.cross_entropy(logits, y, weight=weights) + .6 * F.smooth_l1_loss(reg, r)
            eligible = changed & teacher_eligible[idx]
            if eligible.any():
                selected = torch.as_tensor(eligible, device=device)
                target = torch.as_tensor(teacher_prob[idx][eligible], device=device)
                softened = F.softmax(torch.log(target.clamp_min(1e-8)) / 2, -1)
                loss += .2 * 4 * F.kl_div(F.log_softmax(logits[selected]/2, -1),
                                          softened, reduction='batchmean')
                loss += .1 * F.smooth_l1_loss(reg[selected],
                        torch.as_tensor(teacher_reg[idx][eligible], device=device))
            optimizer.zero_grad(set_to_none=True)
            loss.backward()
            nn.utils.clip_grad_norm_(model.parameters(), 1.)
            optimizer.step()
            losses.append(float(loss.detach()))
        result, _, _ = profile('full_student', model, valid, quick, sparse, device)
        row = {'stage':'student', 'epoch':epoch, 'train_loss':float(np.mean(losses)), **result}
        history.append(row)
        print(json.dumps(row, ensure_ascii=False), flush=True)
        write_csv(out / '完整数据_缺失蒸馏学生_训练记录.csv', history)
        if result['selection_score'] > best + 1e-4:
            best, stale = result['selection_score'], 0
            save_checkpoint(checkpoint, model, 'distilled', seed, epoch, best, stats, config)
        else:
            stale += 1
        if stale >= patience:
            break
    model.load_state_dict(torch.load(checkpoint, map_location=device, weights_only=False)['model'])
    model.eval()
    return model, checkpoint


def target_validation_rows(model, full_valid, target_ids, scenarios_list, device, model_name):
    lookup = {sample_id: i for i, sample_id in enumerate(full_valid['ids'])}
    indices = np.asarray([lookup[x] for x in sorted(target_ids)], dtype=int)
    rows = []
    for scenario in scenarios_list:
        prob, intensity = predict(model, full_valid, device, scenario)
        rows.append({'model':model_name,
                     **{key: scenario[key] for key in ('subset','rate','position','replicate')},
                     **metrics(full_valid['y'][indices], full_valid['r'][indices],
                               prob[indices], intensity[indices])})
    return rows


def main(args):
    runtime(args.device)
    seed_all(args.seed)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    source = Path(args.full_data)
    label_csv = source.with_name('label.csv')
    source_hash = sha(source)
    started = time.time()
    with source.open('rb') as stream:
        data = pickle.load(stream)
    summary, target_ids = audit_splits(data, Path(args.data_root), label_csv,
                                      Path(args.safety_audit))
    train_raw = data['train']
    valid_raw = data['valid']
    data.pop('test', None)  # Never inspect full-test records or labels.
    del data
    gc.collect()
    summary['joint_zero_unk_masked_train'] = hide_joint_unknown(train_raw)
    summary['joint_zero_unk_masked_valid'] = hide_joint_unknown(valid_raw)
    for split in (train_raw, valid_raw):
        split.pop('text', None)  # Unused 768-D source embedding; BERT token IDs are retained.
    dump(out / '完整数据_泄漏与标签审计.json', summary)
    encoder = TextEncoder(args.device)
    cache = ROOT / 'work/问题2_特征缓存'
    train = prepare_split(train_raw, encoder, cache, 'full_train', source_hash)
    valid = prepare_split(valid_raw, encoder, cache, 'full_valid', source_hash)
    del train_raw, valid_raw
    gc.collect()
    stats = fit_normalization(train)
    normalize(train, stats)
    normalize(valid, stats)
    config = {'source_sha256':source_hash, 'label_csv_sha256':sha(label_csv),
              'code_sha256':sha(__file__), 'seed':args.seed, 'bert_model':encoder.revision,
              'python':sys.version.split()[0], 'torch':torch.__version__,
              'train_samples':len(train['ids']), 'validation_samples':len(valid['ids']),
              'architecture':'Fusion hidden96, one temporal Transformer; clean teacher ungated, student gated',
              'teacher_loss':'class-weighted CE + 0.6 SmoothL1',
              'student_loss':'teacher loss + 0.2*4 KL(T=2) + 0.1 SmoothL1 on eligible teacher train predictions',
              'optimizer':'AdamW lr=3e-4 wd=1e-3, batch64, dropout0.2, clip1',
              'teacher_selection':'full valid clean Macro-F1 - 0.1 MAE',
              'student_selection':'0.4 clean F1 + 0.3 contiguous F1 + 0.3 scattered F1 - 0.1 weighted MAE',
              'student_augmentation':'35% clean, 30% contiguous T/A/V subsets 10/30/50%, 35% scattered common TAV 10/20/30%',
              'student_acceptance':'full-valid score >= teacher score + 0.005; clean F1 >= teacher clean F1 - 0.01',
              'test_and_attachment3_policy':'no test or attachment3 data used for fitting, early stopping or selection',
              'attachment2_valid_note':'subset of full valid, descriptive only, not independent selection'}
    dump(out / '完整数据_实验预注册.json', config)
    quick = scenarios(valid, encoder, cache, stats, source_hash)
    sparse = make_scattered_suite(valid, encoder, cache, stats, source_hash)
    teacher, teacher_path = train_teacher(train, valid, quick, args.device, out,
        config, stats, args.seed, args.teacher_epochs, args.patience)
    teacher_profile, teacher_old, teacher_sparse = profile('full_teacher', teacher,
        valid, quick, sparse, args.device)
    write_csv(out / '完整数据_教师验证场景.csv',
              [{'model':'full_teacher', **r} for r in teacher_old+teacher_sparse])
    banks = get_banks(train, encoder, cache, stats, source_hash)
    student, student_path = train_student(train, valid, quick, sparse, teacher, banks,
        args.device, out, config, stats, args.seed, args.student_epochs, args.patience)
    student_profile, student_old, student_sparse = profile('full_student', student,
        valid, quick, sparse, args.device)
    write_csv(out / '完整数据_学生验证场景.csv',
              [{'model':'full_student', **r} for r in student_old+student_sparse])
    accepted = (student_profile['selection_score'] >= teacher_profile['selection_score'] + .005 and
                student_profile['clean_F1'] >= teacher_profile['clean_F1'] - .01)
    selected = 'student' if accepted else 'teacher'
    target_rows = target_validation_rows(teacher, valid, target_ids['valid'],
                                         quick+sparse, args.device, 'full_teacher')
    target_rows += target_validation_rows(student, valid, target_ids['valid'],
                                          quick+sparse, args.device, 'full_student')
    write_csv(out / '附件2验证子集_仅分析.csv', target_rows)
    result = {'selected':selected, 'student_accepted':bool(accepted),
              'teacher_full_valid':teacher_profile, 'student_full_valid':student_profile,
              'teacher_checkpoint':teacher_path.name, 'teacher_sha256':sha(teacher_path),
              'student_checkpoint':student_path.name, 'student_sha256':sha(student_path),
              'elapsed_seconds':time.time()-started,
              'caveat':'Attachment-2 valid is a subset of full valid; attachment2/full test and attachment3 untouched.'}
    dump(out / '完整数据_验证选模结论.json', result)
    print(json.dumps(result, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--full-data', default=ROOT/'work/full_mosei/CMU-MOSEI_完整版_对齐特征与标签/aligned_50.pkl')
    parser.add_argument('--data-root', default=ROOT/'E题数据/E题数据')
    parser.add_argument('--safety-audit', default=ROOT/'outputs/question2/问题2_完整版MOSEI安全整合审计/问题2_完整版MOSEI与竞赛数据安全整合审计.json')
    parser.add_argument('--out', default=ROOT/'outputs/question2/问题2_完整MOSEI训练实验')
    parser.add_argument('--device', default='cuda:1')
    parser.add_argument('--seed', type=int, default=2026)
    parser.add_argument('--teacher-epochs', type=int, default=25)
    parser.add_argument('--student-epochs', type=int, default=30)
    parser.add_argument('--patience', type=int, default=7)
    main(parser.parse_args())
