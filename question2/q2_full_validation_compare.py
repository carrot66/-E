"""第二问：旧三种子集成与完整数据教师/学生的同口径验证对照。

只读取已净化的 train/valid pickle；仅用 valid 标签评价。附件2 label.xlsx
只读取 valid 的样本 ID，用于拆分原728条与新增1143条。绝不访问
test/附件3数据或其标签。各模型沿用自身 checkpoint 的归一化参数和
BERT 修订版，但三种模型共享同一视频分组、同一缺失场景和评分公式。

先运行 q2_prepare_safe_full_data.py 得到 safe_train_valid.pkl，再提供五个
checkpoint。默认权重路径可用命令行覆盖；此脚本不会重新训练或改写权重。
"""
from __future__ import annotations

import argparse
import gc
import json
import pickle
from collections import defaultdict
from pathlib import Path

import numpy as np
import torch
from openpyxl import load_workbook

from q2_ensemble import load_models, predict_ensemble
from q2_experiment import ROOT, TextEncoder, dump, normalize, prepare_split, runtime, scenarios, sha, write_csv
from q2_scattered_gap_experiment import make_scattered_suite, score as declared_score


MODEL_NAMES = ("old_distilled_3seed", "full_clean_teacher", "full_robust_student")
SCENARIO_COUNT = 12  # complete + T/A/V/AV/TAV contiguous 30% + scattered TAV 10/20/30% x2


def video_id(sample_id: str) -> str:
    video, delimiter, clip = str(sample_id).partition("$_$")
    if not delimiter or not video or not clip:
        raise ValueError(f"无法解析视频ID：{sample_id!r}")
    return video


def attachment2_valid_ids(path: Path) -> set[str]:
    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        sheet = workbook["label"]
        result = set()
        for row in sheet.iter_rows(min_row=2):
            # Only inspect split column F. Columns C/D/E contain text/labels and
            # are intentionally not referenced; test rows are skipped entirely.
            if row[5].value != "valid":
                continue
            sid = f"{row[0].value}$_${row[1].value}"
            if sid in result:
                raise ValueError(f"附件2 valid重复ID：{sid}")
            result.add(sid)
        if len(result) != 728:
            raise ValueError(f"附件2 valid应为728条，实际{len(result)}")
        return result
    finally:
        workbook.close()


def checkpoint_revision(checkpoint: dict) -> str:
    provenance = checkpoint.get("provenance", {})
    revision = provenance.get("bert_revision") or provenance.get("bert_model")
    if not revision:
        raise ValueError("checkpoint缺少BERT修订版信息，无法保证输入一致")
    return str(revision)


def load_safe_validation(path: Path) -> dict:
    with path.open("rb") as stream:
        data = pickle.load(stream)
    if set(data) != {"train", "valid"}:
        raise ValueError("输入不是仅含train/valid的safe pickle；拒绝读取完整数据pickle")
    valid = data["valid"]
    if len(valid["id"]) != 1871 or len(data["train"]["id"]) != 16326:
        raise ValueError("safe pickle样本数与独立审计不符")
    del data
    gc.collect()
    if set(valid) != {"id", "raw_text", "text_bert", "audio", "vision",
                      "classification_labels", "regression_labels"}:
        raise ValueError(f"safe valid字段异常：{sorted(valid)}")
    return valid


def prepare_for_checkpoint(raw: dict, checkpoint: dict, encoder: TextEncoder,
                           cache: Path, name: str, source_hash: str) -> dict:
    prepared = prepare_split(raw, encoder, cache, name, source_hash)
    normalize(prepared, checkpoint["stats"])
    if not np.array_equal(prepared["y"], np.sign(prepared["r"]).astype(np.int64) + 1):
        raise ValueError("valid类别与强度符号冲突")
    return prepared


def scenario_signature(scenario: dict) -> tuple:
    return (scenario["subset"], scenario["rate"], scenario["position"], scenario["replicate"])


def make_suite(prepared: dict, encoder: TextEncoder, cache: Path,
               stats: dict, source_hash: str) -> list[dict]:
    quick = scenarios(prepared, encoder, cache, stats, source_hash, full=False)
    scattered = make_scattered_suite(prepared, encoder, cache, stats, source_hash)
    suite = quick + scattered
    if len(suite) != SCENARIO_COUNT:
        raise AssertionError(f"场景数量错误：{len(suite)}")
    expected = [("none", 0., "random", 0)] + [
        (subset, .3, "random", 0) for subset in ("T", "A", "V", "AV", "TAV")
    ] + [("TAV", rate, "scattered", rep) for rate in (.1, .2, .3) for rep in (0, 1)]
    if [scenario_signature(item) for item in suite] != expected:
        raise AssertionError("验证场景次序或定义发生变化")
    return suite


def ensure_paired_inputs(first: dict, second: dict,
                         suite_a: list[dict], suite_b: list[dict]) -> None:
    if first["ids"] != second["ids"]:
        raise ValueError("两组checkpoint的valid样本次序不一致")
    for key in ("y", "r", "valid", "mask"):
        if not np.array_equal(first[key], second[key]):
            raise ValueError(f"两组checkpoint的{key}不一致")
    for a, b in zip(suite_a, suite_b):
        if scenario_signature(a) != scenario_signature(b) or not np.array_equal(a["mask"], b["mask"]):
            raise ValueError("新旧模型的缺失场景并非相同位置")


def macro_f1(y: np.ndarray, pred: np.ndarray) -> float:
    confusion = np.bincount(3*y + pred, minlength=9).reshape(3, 3)
    numerator = 2 * np.diag(confusion)
    denominator = confusion.sum(axis=0) + confusion.sum(axis=1)
    return float(np.mean(np.divide(numerator, denominator,
                                   out=np.zeros(3, dtype=np.float64), where=denominator != 0)))


def scenario_metrics(y: np.ndarray, predicted: np.ndarray, absolute_error: np.ndarray,
                     indices: np.ndarray) -> dict:
    if len(indices) == 0:
        raise ValueError("评价子集为空")
    return {
        "Macro_F1": macro_f1(y[indices], predicted[indices]),
        "MAE": float(np.mean(absolute_error[indices])),
        "Accuracy": float(np.mean(y[indices] == predicted[indices])),
    }


def profile_from_arrays(y: np.ndarray, absolute_error: np.ndarray,
                        predicted: np.ndarray, indices: np.ndarray) -> dict:
    rows = [scenario_metrics(y, predicted[i], absolute_error[i], indices)
            for i in range(SCENARIO_COUNT)]
    score = declared_score(rows[:6], rows[6:])
    return {
        "selection_score": float(score),
        "clean_Macro_F1": rows[0]["Macro_F1"],
        "clean_MAE": rows[0]["MAE"],
        "contiguous_30pct_mean_Macro_F1": float(np.mean([r["Macro_F1"] for r in rows[1:6]])),
        "contiguous_30pct_mean_MAE": float(np.mean([r["MAE"] for r in rows[1:6]])),
        "scattered_TAV_mean_Macro_F1": float(np.mean([r["Macro_F1"] for r in rows[6:]])),
        "scattered_TAV_mean_MAE": float(np.mean([r["MAE"] for r in rows[6:]])),
    }


def grouped_bootstrap_indices(ids: list[str], cohort_indices: np.ndarray,
                              rng: np.random.Generator) -> np.ndarray:
    groups = defaultdict(list)
    for index in cohort_indices:
        groups[video_id(ids[int(index)])].append(int(index))
    videos = sorted(groups)
    draws = rng.integers(0, len(videos), size=len(videos))
    return np.asarray([index for selected in draws for index in groups[videos[int(selected)]]],
                      dtype=np.int64)


def bootstrap_pairs(ids: list[str], y: np.ndarray, errors: dict,
                    predictions: dict, cohorts: dict[str, np.ndarray],
                    replicates: int, seed: int) -> list[dict]:
    pairs = (("full_clean_teacher", "old_distilled_3seed"),
             ("full_robust_student", "old_distilled_3seed"),
             ("full_robust_student", "full_clean_teacher"))
    fields = ("selection_score", "clean_Macro_F1", "clean_MAE",
              "contiguous_30pct_mean_Macro_F1", "contiguous_30pct_mean_MAE",
              "scattered_TAV_mean_Macro_F1", "scattered_TAV_mean_MAE")
    result = []
    for cohort_name, cohort_indices in cohorts.items():
        rng = np.random.default_rng(seed)
        point = {model: profile_from_arrays(y, errors[model], predictions[model], cohort_indices)
                 for model in MODEL_NAMES}
        sampled = {(a, b, field): np.empty(replicates, dtype=np.float64)
                   for a, b in pairs for field in fields}
        for rep in range(replicates):
            take = grouped_bootstrap_indices(ids, cohort_indices, rng)
            profiles = {model: profile_from_arrays(y, errors[model], predictions[model], take)
                        for model in MODEL_NAMES}
            for a, b in pairs:
                for field in fields:
                    sampled[(a, b, field)][rep] = profiles[a][field] - profiles[b][field]
        for a, b in pairs:
            for field in fields:
                values = sampled[(a, b, field)]
                low, high = np.percentile(values, [2.5, 97.5])
                result.append({
                    "cohort": cohort_name, "new_model": a, "reference_model": b,
                    "metric": field, "point_delta_new_minus_reference": point[a][field] - point[b][field],
                    "bootstrap_CI95_low": float(low), "bootstrap_CI95_high": float(high),
                    "bootstrap_fraction_delta_positive": float(np.mean(values > 0)),
                    "bootstrap_unit": "video_id", "bootstrap_replicates": replicates,
                    "bootstrap_seed": seed,
                })
    return result


def main(args: argparse.Namespace) -> None:
    if args.bootstrap_replicates < 100:
        raise ValueError("bootstrap重复次数至少100")
    runtime(args.device)
    old_paths = [Path(path).resolve(strict=True) for path in args.old_checkpoints]
    teacher_path = Path(args.teacher_checkpoint).resolve(strict=True)
    student_path = Path(args.student_checkpoint).resolve(strict=True)
    safe_path = Path(args.safe_data).resolve(strict=True)
    output = Path(args.output)
    output.mkdir(parents=True, exist_ok=True)
    checkpoint_paths = {f"old_distilled_{seed}": path for seed, path in zip((2026, 2027, 2028), old_paths)}
    checkpoint_paths["full_clean_teacher"] = teacher_path
    checkpoint_paths["full_robust_student"] = student_path
    models, checkpoints = load_models(checkpoint_paths, args.device)
    if any(checkpoints[f"old_distilled_{seed}"]["method"] != "distilled" for seed in (2026, 2027, 2028)):
        raise ValueError("旧三权重并非门控蒸馏模型")
    if checkpoints["full_clean_teacher"]["method"] != "baseline" or checkpoints["full_robust_student"]["method"] != "distilled":
        raise ValueError("完整数据教师/学生权重类型错误")
    old = checkpoints["old_distilled_2026"]
    if any(checkpoints[f"old_distilled_{seed}"]["stats"] != old["stats"] or
           checkpoint_revision(checkpoints[f"old_distilled_{seed}"]) != checkpoint_revision(old)
           for seed in (2027, 2028)):
        raise ValueError("旧三权重的BERT版本或归一化参数不一致")
    full = checkpoints["full_clean_teacher"]
    if checkpoints["full_robust_student"]["stats"] != full["stats"] or checkpoint_revision(checkpoints["full_robust_student"]) != checkpoint_revision(full):
        raise ValueError("完整数据教师/学生归一化或BERT版本不一致")
    safe_hash = sha(safe_path)
    safe_manifest_path = safe_path.with_suffix(".manifest.json")
    if safe_manifest_path.exists():
        safe_manifest = json.loads(safe_manifest_path.read_text(encoding="utf-8"))
        if safe_manifest.get("safe_pickle_sha256") != safe_hash:
            raise ValueError("safe pickle与制备清单的SHA-256不一致")
    for seed in (2026, 2027, 2028):
        if int(checkpoints[f"old_distilled_{seed}"].get("seed", -1)) != seed:
            raise ValueError(f"旧checkpoint种子与路径顺序不符：{seed}")
    for name in ("full_clean_teacher", "full_robust_student"):
        provenance = checkpoints[name].get("provenance", {})
        if provenance.get("train_samples") != 16326 or provenance.get("validation_samples") != 1871:
            raise ValueError(f"{name}不是经审计的完整版训练/验证checkpoint")
    raw = load_safe_validation(safe_path)
    old_ids = attachment2_valid_ids(Path(args.attachment2_labels).resolve(strict=True))
    all_ids = list(map(str, raw["id"]))
    if len(set(all_ids)) != 1871 or not old_ids.issubset(all_ids):
        raise ValueError("safe valid与附件2 valid的ID包含关系异常")
    cohorts = {
        "full_valid_1871": np.arange(len(all_ids), dtype=np.int64),
        "attachment2_valid_728": np.flatnonzero(np.isin(all_ids, list(old_ids))),
        "new_full_valid_1143": np.flatnonzero(~np.isin(all_ids, list(old_ids))),
    }
    if [len(value) for value in cohorts.values()] != [1871, 728, 1143]:
        raise AssertionError("验证分层人数不符")
    cache = Path(args.cache)
    encoders = {}
    def encoder_for(checkpoint: dict) -> TextEncoder:
        revision = checkpoint_revision(checkpoint)
        if revision not in encoders:
            encoders[revision] = TextEncoder(args.device, revision)
        return encoders[revision]
    old_encoder, full_encoder = encoder_for(old), encoder_for(full)
    old_valid = prepare_for_checkpoint(raw, old, old_encoder, cache, "full_valid_compare_old", safe_hash)
    full_valid = prepare_for_checkpoint(raw, full, full_encoder, cache, "full_valid_compare_new", safe_hash)
    del raw
    gc.collect()
    old_suite = make_suite(old_valid, old_encoder, cache, old["stats"], safe_hash)
    full_suite = make_suite(full_valid, full_encoder, cache, full["stats"], safe_hash)
    ensure_paired_inputs(old_valid, full_valid, old_suite, full_suite)
    y = old_valid["y"]
    truth = old_valid["r"]
    predicted = {}
    errors = {}
    group_specs = (
        (MODEL_NAMES[0], old_valid, old_suite,
         [f"old_distilled_{seed}" for seed in (2026, 2027, 2028)]),
        (MODEL_NAMES[1], full_valid, full_suite, ["full_clean_teacher"]),
        (MODEL_NAMES[2], full_valid, full_suite, ["full_robust_student"]),
    )
    for name, prepared, suite, members in group_specs:
        classes, absolute = [], []
        for scenario in suite:
            probability, intensity = predict_ensemble(models, members, prepared, args.device, scenario)
            classes.append(probability.argmax(axis=1).astype(np.int8))
            absolute.append(np.abs(truth - intensity).astype(np.float32))
        predicted[name] = np.stack(classes)
        errors[name] = np.stack(absolute)
        print(f"完成{name}在{len(suite)}个共同场景的验证推理", flush=True)
    scenario_rows = []
    profile_rows = []
    for cohort_name, indices in cohorts.items():
        videos = len({video_id(all_ids[int(i)]) for i in indices})
        for model in MODEL_NAMES:
            profile = profile_from_arrays(y, errors[model], predicted[model], indices)
            profile_rows.append({"cohort": cohort_name, "model": model,
                                 "samples": len(indices), "video_groups": videos, **profile})
            for i, scenario in enumerate(old_suite):
                row = scenario_metrics(y, predicted[model][i], errors[model][i], indices)
                scenario_rows.append({"cohort": cohort_name, "model": model,
                                      **{key: scenario[key] for key in ("subset", "rate", "position", "replicate")},
                                      "samples": len(indices), **row})
    intervals = bootstrap_pairs(all_ids, y, errors, predicted, cohorts,
                                args.bootstrap_replicates, args.bootstrap_seed)
    manifest = {
        "safe_pickle_sha256": safe_hash,
        "checkpoint_sha256": {name: sha(path) for name, path in checkpoint_paths.items()},
        "checkpoint_path_names": {name: path.name for name, path in checkpoint_paths.items()},
        "bert_revision": {"old": checkpoint_revision(old), "full": checkpoint_revision(full)},
        "cohort_sizes": {name: len(index) for name, index in cohorts.items()},
        "scenario_definitions": [{key: scenario[key] for key in ("subset", "rate", "position", "replicate")}
                                 for scenario in old_suite],
        "predeclared_score": "0.4 clean Macro-F1 + 0.3 mean five contiguous-30% Macro-F1 + 0.3 mean six scattered-TAV Macro-F1 - 0.1*(same-weight MAE)",
        "prediction_consistency": "All models: Neutral intensity=0, Negative<=-0.01, Positive>=0.01",
        "bootstrap": {"unit": "video_id", "paired": True,
                      "replicates": args.bootstrap_replicates, "seed": args.bootstrap_seed,
                      "CI": "percentile 2.5% and 97.5% of new-minus-reference deltas"},
        "evaluation_scope": "safe pickle valid only; attachment2 xlsx valid IDs only; no test/Annex3 labels or predictions",
        "interpretation_limit": "The full-valid split selected the teacher/student checkpoints; bootstrap CIs describe this fixed selected-model comparison and are not an untouched generalization estimate. Original 728 are part of full valid, not an independent holdout.",
    }
    write_csv(output / "问题2_完整版验证同口径逐场景.csv", scenario_rows)
    write_csv(output / "问题2_完整版验证同口径汇总.csv", profile_rows)
    write_csv(output / "问题2_按视频配对Bootstrap置信区间.csv", intervals)
    dump(output / "问题2_完整版验证同口径对照.json",
         {"manifest": manifest, "summary": profile_rows, "paired_bootstrap": intervals})
    print(json.dumps({"output": str(output), "summary": profile_rows}, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--safe-data", default=ROOT / "work/full_mosei/safe_train_valid.pkl")
    parser.add_argument("--attachment2-labels", default=ROOT / "E题数据/E题数据/附件2-数据集特征文件/label.xlsx")
    parser.add_argument("--old-checkpoints", nargs=3, required=True,
                        metavar=("DISTILLED_2026", "DISTILLED_2027", "DISTILLED_2028"))
    parser.add_argument("--teacher-checkpoint", default=ROOT / "outputs/question2/问题2_完整MOSEI训练实验/完整数据_干净教师_种子2026.pt")
    parser.add_argument("--student-checkpoint", default=ROOT / "outputs/question2/问题2_完整MOSEI训练实验/完整数据_缺失蒸馏学生_种子2026.pt")
    parser.add_argument("--output", default=ROOT / "outputs/question2/问题2_完整版验证同口径对照")
    parser.add_argument("--cache", default=ROOT / "work/问题2_特征缓存")
    parser.add_argument("--device", default="cuda:1")
    parser.add_argument("--bootstrap-replicates", type=int, default=1000)
    parser.add_argument("--bootstrap-seed", type=int, default=20260923)
    main(parser.parse_args())
