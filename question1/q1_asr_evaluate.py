"""独立评估第一问100条音频的ASR WER/CER；不使用参考转写参与解码。"""
from __future__ import annotations

import argparse
import csv
import json
import re
import subprocess
from pathlib import Path

import numpy as np
import openpyxl
import torch
from transformers import AutoModelForCTC, AutoTokenizer

PROJECT_ROOT = Path(__file__).resolve().parents[1]

MODEL = 'facebook/wav2vec2-base-960h'
WORD = re.compile(r"[A-Z0-9]+(?:'[A-Z0-9]+)*")


def normalized_words(text: str) -> list[str]:
    return WORD.findall(text.upper().replace('’', "'"))


def edit_distance(a: list[str], b: list[str]) -> int:
    previous = list(range(len(b) + 1))
    for i, left in enumerate(a, 1):
        current = [i]
        for j, right in enumerate(b, 1):
            current.append(min(current[-1] + 1, previous[j] + 1,
                               previous[j - 1] + (left != right)))
        previous = current
    return previous[-1]


def decode_audio(path: Path) -> np.ndarray:
    cmd = ['ffmpeg', '-nostdin', '-hide_banner', '-loglevel', 'error', '-i', str(path),
           '-vn', '-ac', '1', '-ar', '16000', '-f', 'f32le', 'pipe:1']
    result = subprocess.run(cmd, capture_output=True, check=True)
    wave = np.frombuffer(result.stdout, dtype='<f4').copy()
    if len(wave) < 8000:
        raise ValueError(f'音频不足0.5秒：{path}')
    return np.nan_to_num(wave)


def load_reference(path: Path) -> dict[str, str]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    header = [str(x).strip() if x is not None else '' for x in next(rows)]
    columns = {name: i for i, name in enumerate(header)}
    refs = {}
    def clean_id(value: object) -> str:
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        return str(value).strip()
    for row in rows:
        if row[columns['video_id']] is None or row[columns['clip_id']] is None:
            continue
        sid = f"{clean_id(row[columns['video_id']])}__{clean_id(row[columns['clip_id']])}"
        refs[sid] = str(row[columns['text']] or '')
    workbook.close()
    return refs


@torch.inference_mode()
def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument('--data-root', type=Path, default=PROJECT_ROOT / 'E题数据' / 'E题数据')
    parser.add_argument('--manifest', type=Path, default=PROJECT_ROOT / 'outputs' / '问题1_全量特征结果' / 'sample_manifest.csv')
    parser.add_argument('--out', type=Path, default=PROJECT_ROOT / 'outputs' / '问题1_ASR独立评估')
    parser.add_argument('--device', default='cuda:0')
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    labels = list(args.data_root.rglob('label-100.xlsx'))
    video_roots = [p for p in args.data_root.rglob('MOSEI数据集部分原始视频-100条') if p.is_dir()]
    if len(labels) != 1 or len(video_roots) != 1:
        raise ValueError('标签表或原始视频目录不唯一')
    refs = load_reference(labels[0])
    with args.manifest.open(encoding='utf-8-sig', newline='') as file:
        manifest = list(csv.DictReader(file))
    if len(manifest) != 100 or set(x['sample_id'] for x in manifest) != set(refs):
        raise ValueError('100条清单与标签表不一致')
    device = torch.device(args.device)
    if device.type == 'cuda':
        torch.cuda.set_device(device)
        torch.cuda.set_per_process_memory_fraction(.25, device)
    tokenizer = AutoTokenizer.from_pretrained(MODEL, local_files_only=True)
    model = AutoModelForCTC.from_pretrained(MODEL, local_files_only=True).to(device).eval()
    results = []
    for index, item in enumerate(manifest, 1):
        sid = item['sample_id']
        file = video_roots[0] / item['video_file']
        wave = decode_audio(file)
        rms = float(np.sqrt(np.mean(np.square(wave, dtype=np.float64))))
        silent = rms < 1e-5 or float(np.max(np.abs(wave))) < 1e-4
        # 与第一问提取阶段相同的输入归一化；推理时不提供参考转写。
        normal = (wave - wave.mean()) / np.sqrt(wave.var() + 1e-7)
        tensor = torch.from_numpy(normal.astype(np.float32)).to(device)[None]
        with torch.autocast(device_type=device.type, dtype=torch.float16, enabled=device.type == 'cuda'):
            logits = model(tensor).logits
        ids = logits.argmax(-1).cpu()
        hypothesis = tokenizer.batch_decode(ids)[0]
        ref = refs[sid]
        ref_words, hyp_words = normalized_words(ref), normalized_words(hypothesis)
        ref_chars, hyp_chars = list(' '.join(ref_words)), list(' '.join(hyp_words))
        results.append({'sample_id': sid, 'silent': silent, 'reference': ref,
                        'hypothesis': hypothesis, 'reference_words': len(ref_words),
                        'word_edits': edit_distance(ref_words, hyp_words),
                        'reference_chars': len(ref_chars),
                        'char_edits': edit_distance(ref_chars, hyp_chars),
                        'ctc_alignment': item['ctc_alignment']})
        if index % 10 == 0:
            print(f'ASR {index}/{len(manifest)}', flush=True)
    def summary(rows):
        words = sum(x['reference_words'] for x in rows)
        chars = sum(x['reference_chars'] for x in rows)
        return {'samples': len(rows), 'reference_words': words,
                'word_edits': sum(x['word_edits'] for x in rows),
                'WER': sum(x['word_edits'] for x in rows) / words,
                'reference_chars': chars, 'char_edits': sum(x['char_edits'] for x in rows),
                'CER': sum(x['char_edits'] for x in rows) / chars}
    for row in results:
        row['sample_WER'] = row['word_edits'] / row['reference_words']
        row['sample_CER'] = row['char_edits'] / row['reference_chars']
    methods = sorted(set(x['ctc_alignment'] for x in results))
    result = {'all_100': summary(results),
              'non_silent': summary([x for x in results if not x['silent']]),
              'by_alignment': {method: summary([x for x in results if x['ctc_alignment'] == method]) for method in methods},
              'normalization': 'uppercase; punctuation removed except intraword apostrophe; CER includes spaces',
              'method': 'greedy CTC decode without supplying reference transcript',
              'reference': 'competition label-100.xlsx, not manually corrected',
              'model': MODEL,
              'model_revision': getattr(model.config, '_commit_hash', None),
              'caution': 'alignment groups are selected using confidence from the same CTC model; they are diagnostic subsets, not independent ASR benchmarks'}
    with (args.out / '第一问_ASR逐样本结果.csv').open('w', encoding='utf-8-sig', newline='') as file:
        writer = csv.DictWriter(file, fieldnames=list(results[0]))
        writer.writeheader(); writer.writerows(results)
    review = [x for x in results if x['ctc_alignment'] != 'wav2vec2_ctc_forced_alignment' or x['sample_WER'] >= .8]
    with (args.out / '第一问_ASR人工复核候选.csv').open('w', encoding='utf-8-sig', newline='') as file:
        writer = csv.DictWriter(file, fieldnames=list(results[0]))
        writer.writeheader(); writer.writerows(review)
    (args.out / '第一问_ASR指标.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    main()
